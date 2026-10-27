"""'지역 + ○○크레인' 키워드만 뽑은 별도 엑셀 → keywords/지역_크레인_키워드.xlsx

  python scripts/gen_region_crane_kw.py
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
load = lambda n: json.load(open(os.path.join(ROOT, 'data', f'{n}.json'), encoding='utf-8'))
regions = load('regions')['regions']
page_by_main = {p['main']: p['url'] for p in load('pages')['pages']}

# 지역 단위 이름: 권역 2 + 시군구(키워드명·정식명칭·서울 약칭) + 별칭
BIG = ['크레인', '카고크레인', '하이드로크레인', '미니크레인', '5톤 크레인', '25톤 크레인']
DONG = ['크레인', '카고크레인', '미니크레인']

dongs = {d for r in regions for d in r['dongs']}
rows, seen = [], set()


def add(name, level, r, word):
    kw = f'{name} {word}'
    if kw in seen:
        return
    seen.add(kw)
    # 매핑: 같은 키워드 페이지 → 없으면 상위 지역×장비 → 지역 허브
    eq = word.split()[-1] if '톤' not in word else None
    url = (page_by_main.get(kw)
           or (r and page_by_main.get(f'{r["kw"]} {word}'))
           or (r and eq and page_by_main.get(f'{r["kw"]} {eq}'))
           or (r and page_by_main.get(f'{r["kw"]} 크레인'))
           or page_by_main.get(kw) or '/')
    rows.append([kw, kw.replace(' ', ''), level, r['area'] if r else name.replace(' ', ''), r['name'] if r else '', word, url,
                 '있음' if kw in page_by_main else '상위 페이지'])


for area in ['경기북부', '서울서부']:
    for w in BIG:
        add(area, '권역', None, w)
for r in regions:
    names = [r['kw'], r['name']]
    if r['area'] == '서울서부':
        names.append(r['kw'][:-1])  # 은평구 → 은평
    names += [a for a in r['aliases'] if a not in dongs]
    for i, n in enumerate(dict.fromkeys(names)):
        for w in BIG:
            add(n, '시군구' if i < 2 else '별칭·생활권', r, w)
for r in regions:
    for d in r['dongs']:
        name = f'{r["kw"]} {d}' if d.endswith(('읍', '면')) or d in ('구산동',) else d
        for w in DONG:
            add(name, '동·읍면', r, w)

wb = Workbook()
ws = wb.active
ws.title = '지역+크레인 키워드'
head = ['No', '키워드', '붙여쓰기', '지역 단위', '권역', '상위 지역', '크레인 종류', '매핑 URL', '전용 페이지', '월검색수 PC', '월검색수 모바일', '비고']
ws.append(head)
for c in ws[1]:
    c.font, c.fill, c.alignment = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F3A5F'), Alignment(horizontal='center')
for i, r in enumerate(rows, 1):
    ws.append([i] + r + [None, None, None])
for i, w in enumerate([6, 28, 26, 12, 10, 10, 14, 32, 11, 12, 14, 16], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'A2'
ws.auto_filter.ref = ws.dimensions

s = wb.create_sheet('요약')
from collections import Counter
s.append(['구분', '값', '키워드 수'])
for (k, v), n in sorted(Counter((r[2], r[5]) for r in rows).items()):
    s.append([k, v, n])
s.append(['합계', '', len(rows)])
for i, w in enumerate([14, 16, 10], 1):
    s.column_dimensions[get_column_letter(i)].width = w

out = os.path.join(ROOT, 'keywords', '지역_크레인_키워드.xlsx')
wb.save(out)
sys.stdout.reconfigure(encoding='utf-8')
print(Counter(r[2] for r in rows), len(rows), out)
