"""파워크레인 키워드 엑셀 → data/site-pages.json (사이트 라우트의 단일 출처)

  python -I scripts/build_pages.py [keywords/파워크레인지역_원본.xlsx]

오픈 단계(wave)
  1 = 시군구 (지역 대표 페이지)
  2 = 별칭·생활권
  3 = 동·읍면 × 크레인·카고크레인·미니크레인
  4 = 동·읍면 × 나머지(톤수·하이드로·임대·대여)
"""
import json, os, sys
from collections import Counter
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'keywords', '파워크레인지역_원본.xlsx')
regions = json.load(open(os.path.join(ROOT, 'data', 'regions.json'), encoding='utf-8'))['regions'] + \
          json.load(open(os.path.join(ROOT, 'data', 'regions-incheon.json'), encoding='utf-8'))['regions']
names = {r['name'] for r in regions}

pages = []
for row in load_workbook(SRC, read_only=True).active.iter_rows(min_row=2, values_only=True):
    no, kw, _, level, area, parent, ctype, _ = row[:8]
    if not kw:
        continue
    assert parent in names, f'지역 데이터 없음: {parent} ({kw})'
    assert kw.endswith(ctype), kw
    place = kw[: -len(ctype)].strip()
    if level == '시군구':
        wave = 1
    elif level == '별칭·생활권':
        wave = 2
    else:
        wave = 3 if ctype in ('크레인', '카고크레인', '미니크레인') else 4
    pages.append(dict(no=no, kw=kw, slug=kw.replace(' ', '-'), place=place, type=ctype,
                      level=level, area=area, parent=parent, wave=wave))

# 표기만 다른 같은 장소(강화/강화군, 부평/부평구/인천 부평구, 가정동/인천 가정동) → 엑셀 순서상 첫 페이지를 대표로
import re
norm = lambda place: re.sub(r'(시|군|구)$', '', place.split(' ')[-1])
main = {}
for pg in pages:
    key = (pg['parent'], pg['type'], norm(pg['place']), pg['level'] == '동·읍면')
    if key in main:
        pg['variantOf'] = main[key]
    else:
        main[key] = pg['slug']

slugs = Counter(p['slug'] for p in pages)
assert all(v == 1 for v in slugs.values()), [k for k, v in slugs.items() if v > 1]

json.dump({'pages': pages}, open(os.path.join(ROOT, 'data', 'site-pages.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
sys.stdout.reconfigure(encoding='utf-8')
c = Counter(p['wave'] for p in pages)
acc = 0
for w in sorted(c):
    acc += c[w]
    print(f'wave {w}: {c[w]} (누적 {acc})')
print('표기 변형 페이지(variantOf):', sum(1 for p in pages if 'variantOf' in p))
