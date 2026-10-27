"""축 데이터(data/*.json) → 페이지 목록(data/pages.json) + 키워드 엑셀(keywords/크레인_키워드_페이지기획.xlsx)

  python scripts/gen_pages.py

pages.json 은 나중에 Next.js 라우트(generateStaticParams)가 그대로 읽는 단일 출처다.
"""
import json, os, sys
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
load = lambda n: json.load(open(os.path.join(ROOT, 'data', f'{n}.json'), encoding='utf-8'))
regions, equipment, tons, uses = load('regions')['regions'], load('equipment')['equipment'], load('tons')['tons'], load('uses')['uses']

slug = lambda kw: kw.replace(' ', '-')
pages, keywords = [], []


def add(kind, wave, main, title, desc, h1, region=None, extra=None, subs=()):
    s = slug(main)
    pages.append(dict(kind=kind, wave=wave, slug=s, url=f'/{s}/', main=main, title=title, desc=desc, h1=h1,
                      region=region['slug'] if region else None, area=region['area'] if region else '전체', **(extra or {})))
    seen = set()
    for k, intent in [(main, '메인'), (main.replace(' ', ''), '붙여쓰기')] + [(x, '서브') for x in subs]:
        if k not in seen:
            seen.add(k)
            keywords.append(dict(kw=k, type=intent, kind=kind, wave=wave, url=f'/{s}/', area=region['area'] if region else '전체',
                                 region=region['name'] if region else ''))


# ---- 허브 ----
add('home', 1, '크레인', '경기북부·서울서부 크레인 임대 | 카고·하이드로·스카이차', '고양·파주·김포·의정부부터 은평·마포·강서까지, 카고크레인·하이드로·스카이차·사다리차 당일 배차와 무료 견적.', '경기북부·서울서부 크레인 임대')
pages[-1].update(slug='', url='/')
for area in ['경기북부', '서울서부']:
    names = '·'.join(r['kw'] for r in regions if r['area'] == area)
    add('area', 1, f'{area} 크레인', f'{area} 크레인 임대 | {names[:30]}…', f'{area} {names} 지역 크레인·스카이차 배차 안내와 지역별 견적 기준.', f'{area} 크레인 임대',
        subs=[f'{area} 크레인 임대', f'{area} 스카이차', f'{area} 카고크레인'])
for e in equipment:
    add('equip_hub', 1, f'{e["name"]} 임대', f'{e["name"]} 임대 — {e["tagline"]} | 경기북부·서울서부', f'{e["name"]}({e["range"]}) 임대 가격 기준과 {e["pageQuestion"]} 판단법.', f'{e["name"]} 임대',
        extra=dict(equip=e['slug']), subs=[e['name']] + e['secondary'])
for t in tons:
    add('ton_hub', 1, f'{t["name"]} 크레인', f'{t["name"]} 크레인 임대 — {t["typical"]} | 제원·가격', f'{t["name"]} {t["type"]} 제원({t["reach"]})과 쓰이는 작업, 임대 시 확인할 점.', f'{t["name"]} 크레인 임대',
        extra=dict(ton=t['slug']), subs=t['secondary'])
for u in uses:
    add('use_hub', 1, u['main'], f'{u["main"]} — {u["name"]}에 맞는 장비와 비용', f'{u["name"]}에 쓰는 장비({"·".join(u["equip"])})와 톤수 고르는 법, 작업 전 확인할 점.', f'{u["name"]} {u["main"].split()[-1]}',
        extra=dict(use=u['slug']), subs=u['secondary'])

# ---- 지역 ----
for r in regions:
    k = r['kw']
    add('region', 1, f'{k} 크레인', f'{k} 크레인 임대 | {r["hotspots"][0]}·{r["hotspots"][1]} 당일 배차', f'{r["containedIn"]} 크레인·스카이차·사다리차 배차. {r["demand"][0]}, {r["demand"][1]} 견적 안내.', f'{k} 크레인',
        region=r, subs=[f'{k} 크레인 임대', f'{k} 크레인 가격', f'{k} 크레인 업체', f'{k} 크레인 대여'] + [f'{a} 크레인' for a in r['aliases']])
    for e in equipment:
        add('region_equip', 1, f'{k} {e["name"]}', f'{k} {e["name"]} 임대 | {e["tagline"]}', f'{r["name"]} {e["name"]} 배차 — {r["hotspots"][0]}·{r["hotspots"][1]} 현장, {e["range"]}.', f'{k} {e["name"]}',
            region=r, extra=dict(equip=e['slug']),
            subs=[f'{k} {x}' for x in e['secondary']] + [f'{a} {e["name"]}' for a in r['aliases'][:2]])
    for t in tons:
        add('region_ton', 2, f'{k} {t["name"]} 크레인', f'{k} {t["name"]} 크레인 — {t["typical"]} | 임대·가격', f'{r["name"]} {t["name"]} {t["type"]} 배차. {t["fit"]} 작업, {t["reach"]}.', f'{k} {t["name"]} 크레인',
            region=r, extra=dict(ton=t['slug']), subs=[f'{k} {x}' for x in t['secondary']])
    for u in uses:
        add('region_use', 2, f'{k} {u["main"]}', f'{k} {u["main"]} | {u["name"]} 장비·비용 안내', f'{r["name"]} {u["name"]} 작업 — {"·".join(u["equip"])} 중 현장에 맞는 장비와 {u["ton"]} 기준 견적.', f'{k} {u["name"]} {u["main"].split()[-1]}',
            region=r, extra=dict(use=u['slug']), subs=[f'{k} {x}' for x in u['secondary']])

# ---- 동 (중복 동명·읍면은 지역명 접두) ----
dong_count = Counter(d for r in regions for d in r['dongs'])
for r in regions:
    for d in r['dongs']:
        label = f'{r["kw"]} {d}' if dong_count[d] > 1 or d.endswith(('읍', '면')) else d
        for wave, word, e_slug in [(3, '크레인', None), (4, '스카이차', '스카이차')]:
            add('dong', wave, f'{label} {word}', f'{label} {word} | {r["name"]} {word} 당일 배차', f'{r["containedIn"]} {d} {word} 출동. {r["access"]}.', f'{label} {word}',
                region=r, extra=dict(dong=d, equip=e_slug), subs=[f'{label} {word} 임대', f'{label} {word} 가격'])

# ---- 검증 ----
slugs = [p['slug'] for p in pages]
dup = [s for s, c in Counter(slugs).items() if c > 1]
assert not dup, f'슬러그 중복: {dup}'
assert len(set(p['title'] for p in pages)) == len(pages), '제목 중복'

json.dump({'pages': pages}, open(os.path.join(ROOT, 'data', 'pages.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

# ---- 엑셀 ----
KIND = {'home': '메인', 'area': '권역 허브', 'equip_hub': '장비 허브', 'ton_hub': '톤수 허브', 'use_hub': '용도 허브',
        'region': '지역 허브', 'region_equip': '지역×장비', 'region_ton': '지역×톤수', 'region_use': '지역×용도', 'dong': '동×장비'}
HEAD = PatternFill('solid', fgColor='1F3A5F')
wb = Workbook()


def sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True, color='FFFFFF'), HEAD, Alignment(horizontal='center', vertical='center')
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


ws = wb.active; ws.title = '요약'
cnt = Counter((p['kind'], p['wave']) for p in pages)
rows = [[KIND[k], w, n] for (k, w), n in sorted(cnt.items(), key=lambda x: (x[0][1], list(KIND).index(x[0][0])))]
rows += [['합계', '', len(pages)], [], ['키워드 총계', '', len(keywords)]]
for w in sorted({p['wave'] for p in pages}):
    rows.append([f'웨이브 {w} 누적', '', sum(1 for p in pages if p['wave'] <= w)])
sheet(ws, ['페이지 유형', '배포 웨이브', '페이지 수'], rows, [18, 12, 12])

sheet(wb.create_sheet('페이지목록'),
      ['No', '웨이브', '유형', '권역', '지역', 'URL', '메인 키워드', 'title', 'description', 'H1', '장비', '톤수', '용도', '동'],
      [[i + 1, p['wave'], KIND[p['kind']], p['area'], next((r['name'] for r in regions if r['slug'] == p['region']), ''), p['url'], p['main'], p['title'], p['desc'], p['h1'],
        p.get('equip') or '', p.get('ton') or '', p.get('use') or '', p.get('dong') or ''] for i, p in enumerate(pages)],
      [6, 8, 11, 10, 10, 30, 26, 50, 70, 26, 14, 8, 14, 10])

sheet(wb.create_sheet('키워드'),
      ['No', '키워드', '구분', '매핑 URL', '페이지 유형', '웨이브', '권역', '지역', '월검색수 PC', '월검색수 모바일', '경쟁정도', '비고'],
      [[i + 1, k['kw'], k['type'], k['url'], KIND[k['kind']], k['wave'], k['area'], k['region'], None, None, None, None] for i, k in enumerate(keywords)],
      [7, 30, 9, 32, 11, 8, 10, 10, 12, 14, 10, 20])

sheet(wb.create_sheet('지역·동'), ['권역', '지역', '키워드명', '행정', '수요 현장', '주요 작업', '출동', '동/읍면 수', '동/읍면'],
      [[r['area'], r['name'], r['kw'], r['containedIn'], ' / '.join(r['hotspots']), ' / '.join(r['demand']), r['access'], len(r['dongs']), ', '.join(r['dongs'])] for r in regions],
      [10, 10, 10, 18, 50, 40, 50, 10, 90])

sheet(wb.create_sheet('장비·톤수·용도'), ['축', '이름', '제원/메인키워드', '주요 작업', '서브 키워드'],
      [['장비', e['name'], e['range'], ' / '.join(e['jobs']), ', '.join(e['secondary'])] for e in equipment]
      + [['톤수', t['name'], f'{t["type"]} · {t["reach"]}', t['fit'], ', '.join(t['secondary'])] for t in tons]
      + [['용도', u['name'], u['main'], ' / '.join(u['equip']) + f' (기준 {u["ton"]})', ', '.join(u['secondary'])] for u in uses],
      [8, 16, 40, 50, 60])

os.makedirs(os.path.join(ROOT, 'keywords'), exist_ok=True)
out = os.path.join(ROOT, 'keywords', '크레인_키워드_페이지기획.xlsx')
wb.save(out)
sys.stdout.reconfigure(encoding='utf-8')
for row in rows:
    print(*row)
print('pages', len(pages), 'keywords', len(keywords), '->', out)
