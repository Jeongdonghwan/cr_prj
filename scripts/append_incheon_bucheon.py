"""파워크레인지역.xlsx 에 인천·부천 행을 기존 형식 그대로 이어 붙인다 (원본은 _백업.xlsx 로 복사).

  python -I scripts/append_incheon_bucheon.py "<파워크레인지역.xlsx>"

지역 원본: curtain_prj/data/regions-kr.json (전국 행정동 2026-10 정규화)
단어: 시군구·별칭 12종, 동·읍면 8종 (톤수·임대 포함) / 법정동: data/incheon_bucheon_bdong.json
매핑 URL = 동 '크레인'만 동 페이지, 나머지는 상위 지역 페이지
"""
import json, os, shutil, sys
from collections import Counter
from openpyxl import load_workbook

SRC = sys.argv[1]
REGIONS = json.load(open(os.path.join(os.path.dirname(__file__), '..', '..', 'curtain_prj', 'data', 'regions-kr.json'), encoding='utf-8'))['regions']
BIG = ['크레인', '카고크레인', '하이드로크레인', '미니크레인', '5톤 크레인', '10톤 크레인', '15톤 크레인',
       '25톤 크레인', '50톤 크레인', '100톤 크레인', '크레인 임대', '크레인 대여']
DONG = ['크레인', '카고크레인', '미니크레인', '하이드로크레인', '5톤 크레인', '25톤 크레인', '크레인 임대', '크레인 대여']
AREA = '인천·부천'

# 인천 구·군: 약칭 → (정식명, 페이지 키)
INCHEON_GU = {'제물포': '제물포구', '영종': '영종구', '미추홀': '미추홀구', '연수': '연수구', '남동': '남동구',
              '부평': '부평구', '계양': '계양구', '서해': '서해구', '검단': '검단구', '강화': '강화군', '옹진': '옹진군'}
BUCHEON_GU = {'원미': '원미구', '소사': '소사구', '오정': '오정구'}
# 생활권·옛 구 이름 → 상위 지역
INCHEON_ALIAS = {'인천 중구': '제물포구', '인천 동구': '제물포구', '인천 서구': '서해구', '인천 남구': '미추홀구',
                 '송도': '연수구', '송도국제도시': '연수구', '송도신도시': '연수구', '청라': '서해구', '청라국제도시': '서해구',
                 '루원시티': '서해구', '검단신도시': '검단구', '영종도': '영종구', '인천공항': '영종구', '인천항': '제물포구',
                 '남동공단': '남동구', '남동산단': '남동구', '구월': '남동구', '주안': '미추홀구', '강화도': '강화군'}
# 두 법정동을 묶은 행정동 → 각 법정동 이름도 추가
SPLIT = {'작전서운동': ['서운동'], '논현고잔동': ['고잔동'], '남촌도림동': ['남촌동', '도림동'], '불로대곡동': ['불로동', '대곡동'],
         '신현원창동': ['신현동', '원창동'], '검암경서동': ['검암동', '경서동'], '오류왕길동': ['오류동', '왕길동'], '장수서창동': ['장수동']}

slug = lambda s: s.replace(' ', '-')
page = lambda key, w: f'/{slug(key)}-{slug(w)}/'
national = Counter(r['name'].split()[-1] for r in REGIONS if r['level'] in ('동', '읍면'))

new = []  # [키워드, 붙여쓰기, 지역 단위, 권역, 상위 지역, 크레인 종류, 매핑 URL]


def add(name, level, parent, words, url_key, dong_page=False):
    for w in words:
        kw = f'{name} {w}'
        url = page(name, w) if (dong_page and w == '크레인') else page(url_key, w)
        new.append([kw, kw.replace(' ', ''), level, AREA, parent, w, url])


# 시
add('인천', '시군구', '인천광역시', BIG, '인천')
add('인천시', '시군구', '인천광역시', BIG, '인천')
add('부천', '시군구', '부천시', BIG, '부천')
add('부천시', '시군구', '부천시', BIG, '부천')
# 구·군
for s, full in INCHEON_GU.items():
    add(full, '시군구', full, BIG, full)
    add(s, '시군구', full, BIG, full)
    add(f'인천 {full}', '별칭·생활권', full, BIG, full)
for s, full in BUCHEON_GU.items():
    add(full, '별칭·생활권', '부천시', BIG, '부천')
    add(f'부천 {s}', '별칭·생활권', '부천시', BIG, '부천')
for a, gu in INCHEON_ALIAS.items():
    add(a, '별칭·생활권', gu, BIG, gu)

# 동·읍면
for r in REGIONS:
    if r['level'] not in ('동', '읍면'):
        continue
    if r['sido'] == '인천광역시':
        gu = INCHEON_GU.get(r['parent'])
        if not gu:
            continue
        dong = r['name'].split()[-1]
        # 구 약칭 접두('제물포 송현동')는 검색형인 '인천 송현동'으로. 강화·옹진 읍면은 그대로
        name = r['name'] if ' ' not in r['name'] or r['parent'] in ('강화', '옹진') else f'인천 {dong}'
        add(name, '동·읍면', gu, DONG, gu, dong_page=True)
        for extra in SPLIT.get(dong, []):
            name = extra if national[extra] == 0 else f'인천 {extra}'
            add(name, '동·읍면', gu, DONG, gu, dong_page=True)
    elif r['parent'] in ('부천', *BUCHEON_GU):
        add(r['name'], '동·읍면', '부천시', DONG, '부천', dong_page=True)

# 법정동. 행정동 목록에 없던 동네(운북동·귤현동·여월동 …)는 그대로 넣고,
# 이미 있던 동이나 두 글자 동은 '인천 가정동'·'부천 작동'처럼 시 이름을 붙인 검색형으로 추가된다
BDONG = json.load(open(os.path.join(os.path.dirname(__file__), '..', 'data', 'incheon_bucheon_bdong.json'), encoding='utf-8'))['bdong']
for gu, names in BDONG.items():
    city = '부천' if gu == '부천' else '인천'
    for b in names:
        if b.endswith(('읍', '면')):
            continue  # 읍면은 위 행정동 목록에 이미 있다
        name = b if national[b] == 0 and len(b) > 2 else f'{city} {b}'
        if city == '부천':
            add(name, '동·읍면', '부천시', DONG, '부천', dong_page=True)
        else:
            add(name, '동·읍면', gu, DONG, gu, dong_page=True)

# 중복 제거 + 기존 키워드와 겹치는 것 제외
wb = load_workbook(SRC)
ws = wb.active
existing = {row[1] for row in ws.iter_rows(min_row=2, values_only=True)}
last_no = max(row[0] for row in ws.iter_rows(min_row=2, values_only=True) if isinstance(row[0], int))
seen, rows = set(existing), []
for r in new:
    if r[0] not in seen:
        seen.add(r[0])
        rows.append(r)

backup = SRC.replace('.xlsx', '_백업.xlsx')
if not os.path.exists(backup):
    shutil.copy2(SRC, backup)
for i, r in enumerate(rows, 1):
    ws.append([last_no + i] + r)
if ws.auto_filter.ref:
    ws.auto_filter.ref = ws.dimensions
wb.save(SRC)

sys.stdout.reconfigure(encoding='utf-8')
print('added', len(rows), 'from No', last_no + 1, '| backup', backup)
print(Counter(('부천' if r[4] == '부천시' else '인천', r[2]) for r in rows))
